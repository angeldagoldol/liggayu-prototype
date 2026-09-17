#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import os
import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TITLE = "Implementation of ArrayList in a Student Score Management System"
FLOWCHART_ALT_TEXT = (
    "Flow from student score input through validation, ArrayList storage, "
    "average calculation, high and low traversal, and report display."
)
FLOWCHART_PATH = Path(__file__).resolve().parents[1] / "assets" / "student-score-system-flowchart.png"
NAVY = "0B2E59"
GOLD = "D4AF37"
BLACK = "000000"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build the Chapter I-II student score project paper.")
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--logo", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument(
        "--flowchart-output",
        type=Path,
        default=FLOWCHART_PATH,
        help="Optional flowchart PNG destination; defaults to the tracked assets path.",
    )
    return parser.parse_args()


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def reject_template_output_collision(template: Path, output: Path) -> None:
    template_resolved = template.resolve(strict=True)
    output_resolved = output.resolve(strict=False)
    same_resolved_path = template_resolved == output_resolved
    same_existing_file = output.exists() and os.path.samefile(template, output)
    if same_resolved_path or same_existing_file:
        raise ValueError(
            "Refusing to overwrite source template: output resolves to the same file as the source template."
        )


def set_run_font(run, *, size: float | None = None, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "Arial"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Arial")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")
    run.font.color.rgb = RGBColor(0, 0, 0)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_styles(doc: Document) -> None:
    styles = doc.styles

    normal = styles["normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.first_line_indent = Inches(0.5)
    normal.paragraph_format.space_after = Pt(6)

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    title.font.size = Pt(17)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(10)

    heading_1 = styles["Heading 1"]
    heading_1.font.name = "Arial"
    heading_1._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    heading_1._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    heading_1._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    heading_1.font.size = Pt(17)
    heading_1.font.bold = True
    heading_1.font.color.rgb = RGBColor(0, 0, 0)
    heading_1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading_1.paragraph_format.first_line_indent = Inches(0)
    heading_1.paragraph_format.space_before = Pt(0)
    heading_1.paragraph_format.space_after = Pt(14)
    heading_1.paragraph_format.keep_with_next = True

    heading_2 = styles["Heading 2"]
    heading_2.font.name = "Arial"
    heading_2._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    heading_2._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    heading_2._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    heading_2.font.size = Pt(13)
    heading_2.font.bold = True
    heading_2.font.color.rgb = RGBColor(0, 0, 0)
    heading_2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading_2.paragraph_format.first_line_indent = Inches(0)
    heading_2.paragraph_format.space_before = Pt(10)
    heading_2.paragraph_format.space_after = Pt(7)
    heading_2.paragraph_format.keep_with_next = True

    if "List Number" not in styles:
        styles.add_style("List Number", WD_STYLE_TYPE.PARAGRAPH)
    list_number = styles["List Number"]
    list_number.font.name = "Arial"
    list_number._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    list_number._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    list_number._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    list_number.font.size = Pt(11)
    list_number.font.color.rgb = RGBColor(0, 0, 0)
    list_number.paragraph_format.line_spacing = 1.15
    list_number.paragraph_format.space_after = Pt(5)

    if "Cover Text" not in styles:
        styles.add_style("Cover Text", WD_STYLE_TYPE.PARAGRAPH)
    cover = styles["Cover Text"]
    cover.base_style = normal
    cover.font.name = "Arial"
    cover.font.size = Pt(11)
    cover.font.color.rgb = RGBColor(0, 0, 0)
    cover.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.first_line_indent = Inches(0)
    cover.paragraph_format.line_spacing = 1.0
    cover.paragraph_format.space_after = Pt(3)

    if "TOC Heading" not in styles:
        styles.add_style("TOC Heading", WD_STYLE_TYPE.PARAGRAPH)
    toc_heading = styles["TOC Heading"]
    toc_heading.base_style = normal
    toc_heading.font.name = "Arial"
    toc_heading.font.size = Pt(17)
    toc_heading.font.bold = True
    toc_heading.font.color.rgb = RGBColor(0, 0, 0)
    toc_heading.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc_heading.paragraph_format.first_line_indent = Inches(0)
    toc_heading.paragraph_format.space_after = Pt(16)

    if "TOC Entry" not in styles:
        styles.add_style("TOC Entry", WD_STYLE_TYPE.PARAGRAPH)
    toc_entry = styles["TOC Entry"]
    toc_entry.base_style = normal
    toc_entry.font.name = "Arial"
    toc_entry.font.size = Pt(11)
    toc_entry.font.color.rgb = RGBColor(0, 0, 0)
    toc_entry.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    toc_entry.paragraph_format.first_line_indent = Inches(0)
    toc_entry.paragraph_format.line_spacing = 1.15
    toc_entry.paragraph_format.space_after = Pt(7)
    toc_entry.paragraph_format.tab_stops.add_tab_stop(
        Inches(6.2), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )

    if "Specific Objective" not in styles:
        styles.add_style("Specific Objective", WD_STYLE_TYPE.PARAGRAPH)
    objective = styles["Specific Objective"]
    objective.base_style = normal
    objective.font.name = "Arial"
    objective.font.size = Pt(11)
    objective.font.color.rgb = RGBColor(0, 0, 0)
    objective.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    objective.paragraph_format.left_indent = Inches(0.35)
    objective.paragraph_format.first_line_indent = Inches(-0.25)
    objective.paragraph_format.line_spacing = 1.15
    objective.paragraph_format.space_after = Pt(4)

    if "Figure Caption" not in styles:
        styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption = styles["Figure Caption"]
    caption.base_style = normal
    caption.font.name = "Arial"
    caption.font.size = Pt(10)
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Inches(0)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(6)
    caption.paragraph_format.keep_with_next = True


def clear_template_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_section(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)
    for pg_num in section._sectPr.findall(qn("w:pgNumType")):
        section._sectPr.remove(pg_num)


def clear_story(story) -> None:
    for paragraph in story.paragraphs:
        for child in list(paragraph._p):
            paragraph._p.remove(child)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, size=10)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, text, end))


def restart_page_numbering(section) -> None:
    pg_num = OxmlElement("w:pgNumType")
    pg_num.set(qn("w:fmt"), "decimal")
    pg_num.set(qn("w:start"), "1")
    section._sectPr.append(pg_num)


def set_alt_text(inline_shape, description: str, name: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", description)
    doc_pr.set("name", name)


def add_centered_picture(doc: Document, path: Path, width: float, alt_text: str, name: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Inches(0)
    paragraph.paragraph_format.space_after = Pt(3)
    inline_shape = paragraph.add_run().add_picture(str(path), width=Inches(width))
    set_alt_text(inline_shape, alt_text, name)
    return paragraph


def add_cover(doc: Document, logo: Path) -> None:
    picture = add_centered_picture(
        doc,
        logo,
        1.35,
        "Seal of St. John Paul II College of Davao.",
        "SJPIICD Logo",
    )
    picture.paragraph_format.space_after = Pt(7)

    institution = doc.add_paragraph("ST. JOHN PAUL II COLLEGE OF DAVAO", style="Cover Text")
    set_run_font(institution.runs[0], size=13, bold=True)
    course = doc.add_paragraph("Data Structures and Algorithms", style="Cover Text")
    set_run_font(course.runs[0], size=11)

    project = doc.add_paragraph("MINI-SYSTEM PROJECT", style="Cover Text")
    project.paragraph_format.space_before = Pt(8)
    project.paragraph_format.space_after = Pt(5)
    set_run_font(project.runs[0], size=12, bold=True)

    title = doc.add_paragraph(TITLE, style="Title")
    title.paragraph_format.keep_with_next = True

    cover_fields = (
        ("Researchers", "[Full Names of Group Members]"),
        ("Group Leader", "[Full Name]"),
        ("Instructor", "[Instructor's Name]"),
        ("Section", "[Section Name/Code]"),
        ("Date of Submission", "[Month, Day, Year]"),
    )
    for label, value in cover_fields:
        label_paragraph = doc.add_paragraph(label, style="Cover Text")
        label_paragraph.paragraph_format.space_before = Pt(3)
        set_run_font(label_paragraph.runs[0], size=10, bold=True)
        value_paragraph = doc.add_paragraph(value, style="Cover Text")
        set_run_font(value_paragraph.runs[0], size=11)


def add_toc_entry(doc: Document, label: str, page: int, *, subsection: bool = False) -> None:
    paragraph = doc.add_paragraph(style="TOC Entry")
    if subsection:
        paragraph.paragraph_format.left_indent = Inches(0.3)
    paragraph.add_run(f"{label}\t{page}")


def add_static_toc(doc: Document, pages: dict[str, int]) -> None:
    heading = doc.add_paragraph("TABLE OF CONTENTS", style="TOC Heading")
    heading.paragraph_format.space_before = Pt(3)
    entries = (
        ("Chapter I Introduction", "chapter1", False),
        ("1.1 Background of the Study", "background", True),
        ("1.2 Objectives of the Study", "objectives", True),
        ("1.3 Significance of the Study", "significance", True),
        ("Chapter II System Overview and Design", "chapter2", False),
        ("2.1 System Description", "description", True),
        ("2.2 System Architecture and Flowchart", "architecture", True),
        ("2.3 System Features", "features", True),
    )
    for label, key, subsection in entries:
        add_toc_entry(doc, label, pages[key], subsection=subsection)


def add_body_paragraph(doc: Document, text: str, *, lead: str | None = None):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Inches(0.5)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(5)
    if lead:
        lead_run = paragraph.add_run(lead)
        set_run_font(lead_run, size=11, bold=True)
    run = paragraph.add_run(text)
    set_run_font(run, size=11)
    return paragraph


def add_manual_numbered_paragraph(doc: Document, number: int, text: str) -> None:
    paragraph = doc.add_paragraph(style="Specific Objective")
    run = paragraph.add_run(f"{number}.  {text}")
    set_run_font(run, size=11)


def add_chapter_one(doc: Document) -> None:
    doc.add_paragraph("CHAPTER I INTRODUCTION", style="Heading 1")

    doc.add_paragraph("1.1 Background of the Study", style="Heading 2")
    add_body_paragraph(
        doc,
        "Recording student scores requires a structure that can hold a changing number of records during a class session. "
        "A fixed-size arrangement would require a capacity decision before the number of entries is known. A Java ArrayList "
        "allows the system to append each accepted student record as it is submitted while preserving insertion order.",
    )
    add_body_paragraph(
        doc,
        "Each stored Student object contains a name, Assessment 1, Assessment 2, and Assessment 3. Duplicate names are "
        "allowed because different students may share the same name. The application calculates the arithmetic mean of the "
        "three scores and rounds the reported average to two decimal places in Java.",
    )
    add_body_paragraph(
        doc,
        "The prototype connects this data structure to a small web application. Java validates scores from 0 through 100, "
        "stores valid records in ArrayList<Student>, and traverses the list to determine the highest and lowest reported "
        "averages, including every tied name. Records remain in memory only and reset when the Java process restarts.",
    )

    doc.add_paragraph("1.2 Objectives of the Study", style="Heading 2")
    add_body_paragraph(
        doc,
        "The general objective of the project is to implement ArrayList in a working Student Score Management System and "
        "observe how the data structure supports record storage and report generation.",
    )
    objectives = (
        "Store a flexible number of Student objects in a Java ArrayList while preserving their order of entry.",
        "Accept a student name and three assessment scores, then reject blank, nonnumeric, or out-of-range values.",
        "Calculate each student's arithmetic mean in Java and report it to two decimal places.",
        "Traverse the stored records to identify all names tied for the highest or lowest reported average.",
        "Present the returned records and summaries in a responsive browser interface and note improvements from the observed workflow.",
    )
    for number, objective in enumerate(objectives, start=1):
        add_manual_numbered_paragraph(doc, number, objective)

    doc.add_paragraph("1.3 Significance of the Study", style="Heading 2")
    add_body_paragraph(
        doc,
        " The project connects ArrayList operations with input validation, average calculation, and list traversal.",
        lead="Researchers.",
    )
    add_body_paragraph(
        doc,
        " Students can trace each record from form input to dynamic-list storage and report output.",
        lead="Data Structures and Algorithms students.",
    )
    add_body_paragraph(
        doc,
        " Instructors can demonstrate ArrayList storage, validation, and tie handling with a working web example.",
        lead="Instructors.",
    )
    add_body_paragraph(
        doc,
        " Beginners can study a focused Java collection without persistent storage or extra record-management functions.",
        lead="Beginner programmers.",
    )


def add_chapter_two(doc: Document, flowchart: Path) -> None:
    doc.add_paragraph("CHAPTER II SYSTEM OVERVIEW AND DESIGN", style="Heading 1")

    doc.add_paragraph("2.1 System Description", style="Heading 2")
    add_body_paragraph(
        doc,
        "The Student Score Management System runs as one Spring Boot service. Spring Boot serves the HTML, CSS, and JavaScript "
        "interface and exposes REST endpoints for adding a student and retrieving the current report. A user enters a student "
        "name with Assessment 1, Assessment 2, and Assessment 3 through the browser form.",
    )
    add_body_paragraph(
        doc,
        "The browser performs required-field checks and sends the values to Java. Java applies the authoritative validation, "
        "creates a Student object, appends it to the in-memory ArrayList, calculates averages, and traverses the records for "
        "the highest and lowest results. The browser formats and displays the returned values but does not calculate averages "
        "or select the highest and lowest students.",
    )
    add_body_paragraph(
        doc,
        "The service returns student rows in insertion order together with the highest and lowest summaries. The collection is "
        "kept in memory, so the current records reset whenever the Java process restarts or the deployed service is redeployed.",
    )

    doc.add_paragraph("2.2 System Architecture and Flowchart", style="Heading 2")
    add_body_paragraph(
        doc,
        "The architecture keeps the browser responsible for data entry and presentation while Java owns validation, storage, "
        "average calculation, and high and low traversal. Figure 1 shows the implemented request and response flow, including "
        "the recoverable error path that retains the form values.",
    )
    image_paragraph = add_centered_picture(
        doc,
        flowchart,
        4.65,
        FLOWCHART_ALT_TEXT,
        "Student Score Management System Flowchart",
    )
    image_paragraph.paragraph_format.keep_with_next = True
    caption = doc.add_paragraph("Figure 1 Student Score Management System Flowchart", style="Figure Caption")
    caption.paragraph_format.keep_with_next = False
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    add_body_paragraph(
        doc,
        " The input consists of the student's name and the three assessment scores entered in the browser. Each score must be "
        "numeric and fall from 0 through 100 inclusive.",
        lead="Input.",
    )
    add_body_paragraph(
        doc,
        " The browser checks required fields before submission. Java then validates the request, creates the Student object, "
        "appends it to ArrayList<Student>, calculates rounded averages, and traverses the list to collect all names tied for "
        "the highest and lowest values.",
        lead="Process.",
    )
    add_body_paragraph(
        doc,
        " The Java service returns JSON containing the student rows and the high and low summaries. JavaScript renders the "
        "formatted table and summaries, then leaves the interface ready for another entry.",
        lead="Output.",
    )

    doc.add_paragraph("2.3 System Features", style="Heading 2")
    features = (
        "Accepts a student name and exactly three assessment scores through a responsive browser form.",
        "Validates required fields and scores from 0 through 100, while retaining entered values when submission fails.",
        "Stores each accepted Student object in ArrayList<Student> in insertion order and allows duplicate student names.",
        "Calculates the arithmetic mean of the three scores in Java and reports the average to two decimal places.",
        "Traverses the stored records in Java to determine the highest and lowest reported averages.",
        "Returns every student name tied for the highest or lowest average.",
        "Displays the three scores, calculated average, and high and low summaries without reloading the page.",
        "Handles an empty collection as a normal state and keeps accepted records in memory until the Java process restarts.",
    )
    for number, feature in enumerate(features, start=1):
        paragraph = doc.add_paragraph(f"{number}.  {feature}", style="List Number")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.keep_together = True


def rounded_text(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], text: str, font, *, fill: str, outline: str, radius: int = 32) -> None:
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=7)
    x1, y1, x2, y2 = box
    max_width = x2 - x1 - 90
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if draw.textbbox((0, 0), candidate, font=font)[2] <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    text_block = "\n".join(lines)
    spacing = 10
    bbox = draw.multiline_textbbox(
        (0, 0),
        text_block,
        font=font,
        spacing=spacing,
        align="center",
    )
    text_width = bbox[2] - bbox[0]
    text_height = bbox[3] - bbox[1]
    x = (x1 + x2 - text_width) / 2 - bbox[0]
    y = (y1 + y2 - text_height) / 2 - bbox[1]
    draw.multiline_text(
        (x, y),
        text_block,
        font=font,
        fill="#111827",
        spacing=spacing,
        align="center",
    )


def draw_arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], *, color: str = "#0B2E59", width: int = 10) -> None:
    draw.line((start, end), fill=color, width=width)
    x, y = end
    draw.polygon(((x, y), (x - 20, y - 30), (x + 20, y - 30)), fill=color)


def draw_polyline_arrow(draw: ImageDraw.ImageDraw, points: list[tuple[int, int]], *, color: str = "#0B2E59", width: int = 10) -> None:
    draw.line(points, fill=color, width=width, joint="curve")
    x, y = points[-1]
    previous_x, previous_y = points[-2]
    if y > previous_y:
        arrow = ((x, y), (x - 20, y - 30), (x + 20, y - 30))
    elif y < previous_y:
        arrow = ((x, y), (x - 20, y + 30), (x + 20, y + 30))
    elif x > previous_x:
        arrow = ((x, y), (x - 30, y - 20), (x - 30, y + 20))
    else:
        arrow = ((x, y), (x + 30, y - 20), (x + 30, y + 20))
    draw.polygon(arrow, fill=color)


def load_chart_font(size: int, *, bold: bool = False):
    candidates = (
        "/usr/share/fonts/opentype/urw-base35/NimbusSans-Bold.otf" if bold else "/usr/share/fonts/opentype/urw-base35/NimbusSans-Regular.otf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    )
    for candidate in candidates:
        path = Path(candidate)
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def build_flowchart(path: Path) -> None:
    width, height = 1800, 3200
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font = load_chart_font(50)
    bold_font = load_chart_font(52, bold=True)
    small_font = load_chart_font(38, bold=True)

    main_left, main_right = 220, 1270
    node_height, gap = 150, 85
    y = 75
    nodes: list[tuple[str, tuple[int, int, int, int], str, str]] = []
    labels = (
        ("Start", "terminator"),
        ("Enter student name and three scores", "input"),
        ("Browser required-field check", "process"),
        ("Submit request to Java", "process"),
        ("Java validates name and scores", "process"),
        ("Valid input?", "decision"),
        ("Create Student and append to ArrayList<Student>", "process"),
        ("Calculate averages to two decimal places", "process"),
        ("Traverse records for high and low ties", "process"),
        ("Return report as JSON", "output"),
        ("Render table and summaries", "output"),
        ("Await next entry", "terminator"),
    )
    for label, kind in labels:
        h = 190 if kind == "decision" else node_height
        box = (main_left, y, main_right, y + h)
        nodes.append((label, box, kind, GOLD if kind in {"decision", "terminator"} else "DCE8F5"))
        y += h + gap

    for index, (label, box, kind, fill) in enumerate(nodes):
        x1, y1, x2, y2 = box
        if kind == "decision":
            cx, cy = (x1 + x2) // 2, (y1 + y2) // 2
            points = ((cx, y1), (x2, cy), (cx, y2), (x1, cy))
            draw.polygon(points, fill=f"#{fill}", outline=f"#{NAVY}")
            draw.line(points + (points[0],), fill=f"#{NAVY}", width=7, joint="curve")
            bbox = draw.textbbox((0, 0), label, font=bold_font)
            draw.text((cx - (bbox[2] - bbox[0]) / 2, cy - (bbox[3] - bbox[1]) / 2 - 7), label, font=bold_font, fill="#111827")
        else:
            rounded_text(
                draw,
                box,
                label,
                bold_font if kind == "terminator" else font,
                fill=f"#{fill}",
                outline=f"#{NAVY}",
                radius=75 if kind == "terminator" else 30,
            )
        if index < len(nodes) - 1 and index != 5:
            next_box = nodes[index + 1][1]
            draw_arrow(draw, ((x1 + x2) // 2, y2), ((next_box[0] + next_box[2]) // 2, next_box[1]))

    decision = nodes[5][1]
    yes_target = nodes[6][1]
    draw_arrow(
        draw,
        ((decision[0] + decision[2]) // 2, decision[3]),
        ((yes_target[0] + yes_target[2]) // 2, yes_target[1]),
    )
    draw.text((760, decision[3] + 18), "YES", font=small_font, fill=f"#{NAVY}")

    error_box = (1340, decision[1] - 30, 1760, decision[3] + 30)
    rounded_text(draw, error_box, "Show error and retain values", font, fill="#FCE8E6", outline="#A63D40", radius=30)
    draw_polyline_arrow(
        draw,
        [(decision[2], (decision[1] + decision[3]) // 2), (1320, (decision[1] + decision[3]) // 2), (1340, (decision[1] + decision[3]) // 2)],
        color="#A63D40",
    )
    draw.text((1205, decision[1] + 38), "NO", font=small_font, fill="#A63D40")

    input_box = nodes[1][1]
    draw_polyline_arrow(
        draw,
        [
            ((error_box[0] + error_box[2]) // 2, error_box[1]),
            ((error_box[0] + error_box[2]) // 2, input_box[1] - 42),
            (main_right + 45, input_box[1] - 42),
            (main_right + 45, (input_box[1] + input_box[3]) // 2),
            (main_right, (input_box[1] + input_box[3]) // 2),
        ],
        color="#A63D40",
    )

    await_box = nodes[-1][1]
    draw_polyline_arrow(
        draw,
        [
            (await_box[0], (await_box[1] + await_box[3]) // 2),
            (100, (await_box[1] + await_box[3]) // 2),
            (100, (input_box[1] + input_box[3]) // 2),
            (input_box[0], (input_box[1] + input_box[3]) // 2),
        ],
        color=f"#{GOLD}",
        width=12,
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, format="PNG", optimize=True, dpi=(300, 300))


def build_document(
    template: Path,
    logo: Path,
    output: Path,
    flowchart_output: Path | None = None,
) -> None:
    if not template.is_file():
        raise FileNotFoundError(f"Template not found: {template}")
    reject_template_output_collision(template, output)
    flowchart_path = flowchart_output if flowchart_output is not None else FLOWCHART_PATH
    reject_template_output_collision(template, flowchart_path)
    if not logo.is_file():
        raise FileNotFoundError(f"Logo not found: {logo}")

    template_digest = sha256(template)
    build_flowchart(flowchart_path)

    doc = Document(template)
    clear_template_body(doc)
    configure_styles(doc)
    configure_section(doc.sections[0])
    doc.sections[0].footer.is_linked_to_previous = False
    clear_story(doc.sections[0].footer)

    doc.core_properties.title = TITLE
    doc.core_properties.subject = "Data Structures and Algorithms Mini-System Project Chapters I and II"
    doc.core_properties.keywords = "ArrayList, student scores, Java, Spring Boot"

    add_cover(doc, logo)

    content_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(content_section)
    content_section.header.is_linked_to_previous = False
    content_section.footer.is_linked_to_previous = False
    clear_story(content_section.header)
    clear_story(content_section.footer)
    restart_page_numbering(content_section)
    add_page_field(content_section.footer.paragraphs[0])

    toc_pages = {
        "chapter1": 2,
        "background": 2,
        "objectives": 2,
        "significance": 2,
        "chapter2": 3,
        "description": 3,
        "architecture": 3,
        "features": 5,
    }
    add_static_toc(doc, toc_pages)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_chapter_one(doc)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_chapter_two(doc, flowchart_path)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    if sha256(template) != template_digest:
        raise RuntimeError("Source template changed during document generation")


def main() -> None:
    args = parse_args()
    flowchart_output = args.flowchart_output.resolve()
    build_document(
        args.template.resolve(),
        args.logo.resolve(),
        args.output.resolve(),
        flowchart_output,
    )
    print(f"Created {args.output}")
    print(f"Created {flowchart_output}")


if __name__ == "__main__":
    main()
