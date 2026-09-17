from __future__ import annotations

import hashlib
import importlib.util
import os
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "upload" / "DATA STRUCTURES AND ALGORITHMS MINI-SYSTEM PROJECT(format paper presentation) (1).docx"
OUTPUT = ROOT / "deliverables" / "Implementation_of_ArrayList_in_a_Student_Score_Management_System_Chapters_I_II.docx"
TRACKED_FLOWCHART = ROOT / "assets" / "student-score-system-flowchart.png"
EXPECTED_TEMPLATE_SHA256 = "7244151d4da792033142730c69d315e3fb8fb91f978fd0104496aceb0fb61642"
TITLE = "Implementation of ArrayList in a Student Score Management System"
SUBSECTION_HEADINGS = (
    "1.1 Background of the Study",
    "1.2 Objectives of the Study",
    "1.3 Significance of the Study",
    "2.1 System Description",
    "2.2 System Architecture and Flowchart",
    "2.3 System Features",
)
PLACEHOLDERS = (
    "[Full Names of Group Members]",
    "[Full Name]",
    "[Instructor's Name]",
    "[Section Name/Code]",
    "[Month, Day, Year]",
)
FLOWCHART_ALT_TEXT = (
    "Flow from student score input through validation, ArrayList storage, "
    "average calculation, high and low traversal, and report display."
)


@pytest.fixture
def builder_module(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    builder_path = ROOT / "scripts" / "build_chapters_1_2_docx.py"
    spec = importlib.util.spec_from_file_location("task6_docx_builder", builder_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    isolated_flowchart = tmp_path / "isolated-flowchart.png"
    monkeypatch.setattr(module, "FLOWCHART_PATH", isolated_flowchart)
    return module, isolated_flowchart


def _temporary_template(tmp_path: Path) -> Path:
    template = tmp_path / "template.docx"
    shutil.copyfile(TEMPLATE, template)
    return template


def _assert_collision_rejected_without_mutation(builder_module, template: Path, output: Path) -> None:
    module, isolated_flowchart = builder_module
    before = template.read_bytes()

    with pytest.raises(ValueError, match="source template"):
        module.build_document(template, ROOT / "assets" / "sjpiicd-logo.png", output)

    assert template.read_bytes() == before
    assert not isolated_flowchart.exists()


def test_builder_rejects_template_as_output_before_writing(builder_module, tmp_path: Path) -> None:
    template = _temporary_template(tmp_path)
    _assert_collision_rejected_without_mutation(builder_module, template, template)


def test_builder_rejects_symlink_alias_of_template_before_writing(builder_module, tmp_path: Path) -> None:
    template = _temporary_template(tmp_path)
    alias = tmp_path / "template-alias.docx"
    alias.symlink_to(template.name)
    _assert_collision_rejected_without_mutation(builder_module, template, alias)


def test_builder_rejects_hardlink_alias_of_template_before_writing(builder_module, tmp_path: Path) -> None:
    template = _temporary_template(tmp_path)
    alias = tmp_path / "template-hardlink.docx"
    os.link(template, alias)
    _assert_collision_rejected_without_mutation(builder_module, template, alias)


def test_builder_rejects_flowchart_output_matching_template_before_writing(
    builder_module,
    tmp_path: Path,
) -> None:
    module, _ = builder_module
    template = _temporary_template(tmp_path)
    document_output = tmp_path / "chapters.docx"
    before = template.read_bytes()

    with pytest.raises(ValueError, match="source template"):
        module.build_document(
            template,
            ROOT / "assets" / "sjpiicd-logo.png",
            document_output,
            template,
        )

    assert template.read_bytes() == before
    assert not document_output.exists()


@pytest.fixture(scope="session")
def generated_artifacts(tmp_path_factory: pytest.TempPathFactory) -> tuple[Path, Path]:
    build_dir = tmp_path_factory.mktemp("generated-chapters-document")
    output = build_dir / "chapters.docx"
    flowchart = build_dir / "student-score-system-flowchart.png"
    before = hashlib.sha256(TEMPLATE.read_bytes()).hexdigest()
    tracked_flowchart_before = hashlib.sha256(TRACKED_FLOWCHART.read_bytes()).hexdigest()
    result = subprocess.run(
        [
            sys.executable,
            str(ROOT / "scripts" / "build_chapters_1_2_docx.py"),
            "--template",
            str(TEMPLATE),
            "--logo",
            str(ROOT / "assets" / "sjpiicd-logo.png"),
            "--output",
            str(output),
            "--flowchart-output",
            str(flowchart),
        ],
        capture_output=True,
        text=True,
        check=False,
        cwd=build_dir,
    )

    assert result.returncode == 0, result.stderr
    assert output.is_file()
    assert flowchart.is_file()
    assert hashlib.sha256(TEMPLATE.read_bytes()).hexdigest() == before
    assert hashlib.sha256(TRACKED_FLOWCHART.read_bytes()).hexdigest() == tracked_flowchart_before
    return output, flowchart


@pytest.fixture(params=("generated", "delivered"), ids=("generated", "delivered"))
def artifact_path(request: pytest.FixtureRequest, generated_artifacts: tuple[Path, Path]) -> Path:
    return generated_artifacts[0] if request.param == "generated" else OUTPUT


def test_builder_cli_creates_isolated_artifacts_without_modifying_sources(
    generated_artifacts: tuple[Path, Path],
) -> None:
    generated_document, generated_flowchart = generated_artifacts
    assert generated_document.is_file()
    assert generated_flowchart.is_file()


def _document(path: Path) -> Document:
    assert path.exists(), f"Expected document does not exist: {path}"
    return Document(path)


def _paragraph_texts(doc: Document) -> list[str]:
    return [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]


def _xml_parts(path: Path) -> dict[str, bytes]:
    assert path.exists(), f"Expected document does not exist: {path}"
    with zipfile.ZipFile(path) as package:
        return {name: package.read(name) for name in package.namelist()}


def _paragraph_after(doc: Document, heading: str):
    paragraphs = doc.paragraphs
    index = next(i for i, paragraph in enumerate(paragraphs) if paragraph.text.strip() == heading)
    return next(paragraph for paragraph in paragraphs[index + 1 :] if paragraph.text.strip())


def test_output_document_exists() -> None:
    assert OUTPUT.exists(), f"Expected document does not exist: {OUTPUT}"


def test_source_template_checksum_is_unchanged() -> None:
    assert hashlib.sha256(TEMPLATE.read_bytes()).hexdigest() == EXPECTED_TEMPLATE_SHA256


def test_temporary_generated_document_has_chapters_layout_images_and_page_field(
    generated_artifacts: tuple[Path, Path],
) -> None:
    generated_document, _ = generated_artifacts
    doc = _document(generated_document)
    texts = _paragraph_texts(doc)
    parts = _xml_parts(generated_document)

    assert texts.count("CHAPTER I INTRODUCTION") == 1
    assert texts.count("CHAPTER II SYSTEM OVERVIEW AND DESIGN") == 1
    assert all(texts.count(heading) == 1 for heading in SUBSECTION_HEADINGS)
    assert len(doc.sections) == 2
    assert all(section.page_width.inches == pytest.approx(8.5, abs=0.01) for section in doc.sections)
    assert all(section.page_height.inches == pytest.approx(11.0, abs=0.01) for section in doc.sections)
    assert len(doc.inline_shapes) == 2
    document_xml = parts["word/document.xml"].decode("utf-8")
    assert document_xml.count("<wp:inline") == 2
    assert FLOWCHART_ALT_TEXT in document_xml
    footer_xml = "\n".join(
        content.decode("utf-8")
        for name, content in parts.items()
        if re.fullmatch(r"word/footer\d+\.xml", name)
    )
    assert footer_xml.count(" PAGE ") == 1


def test_flowchart_error_message_has_clear_vertical_padding(
    generated_artifacts: tuple[Path, Path],
) -> None:
    _, generated_flowchart = generated_artifacts
    image = Image.open(generated_flowchart).convert("RGB")
    search_left = 1250
    pink = (252, 232, 230)
    pink_points = [
        (x, y)
        for y in range(image.height)
        for x in range(search_left, image.width)
        if image.getpixel((x, y)) == pink
    ]
    assert pink_points
    left = min(x for x, _ in pink_points)
    top = min(y for _, y in pink_points)
    right = max(x for x, _ in pink_points)
    bottom = max(y for _, y in pink_points)
    dark_points = []
    for y in range(top, bottom + 1):
        for x in range(left, right + 1):
            red, green, blue = image.getpixel((x, y))
            if red < 80 and green < 80 and blue < 80:
                dark_points.append((x, y))
    assert dark_points

    top_padding = min(y for _, y in dark_points) - top
    bottom_padding = bottom - max(y for _, y in dark_points)
    assert top_padding >= 28
    assert bottom_padding >= 28


def test_required_title_headings_and_placeholders_are_exact_and_unique(artifact_path: Path) -> None:
    doc = _document(artifact_path)
    texts = _paragraph_texts(doc)

    assert texts.count(TITLE) == 1
    for heading in SUBSECTION_HEADINGS:
        assert texts.count(heading) == 1
        paragraph = next(p for p in doc.paragraphs if p.text.strip() == heading)
        assert paragraph.style.name == "Heading 2"

    chapter_headings = (
        "CHAPTER I INTRODUCTION",
        "CHAPTER II SYSTEM OVERVIEW AND DESIGN",
    )
    for heading in chapter_headings:
        assert texts.count(heading) == 1
        paragraph = next(p for p in doc.paragraphs if p.text.strip() == heading)
        assert paragraph.style.name == "Heading 1"

    for placeholder in PLACEHOLDERS:
        assert texts.count(placeholder) == 1


def test_document_stops_after_chapter_two_and_contains_no_template_instructions(artifact_path: Path) -> None:
    full_text = "\n".join(_paragraph_texts(_document(artifact_path)))
    banned_patterns = (
        r"\bchapter\s+(iii|iv|v|vi)\b",
        r"\breferences\b",
        r"\bappendices\b",
        r"bubble sort",
        r"quick sort",
        r"include the sjpiicd logo",
        r"explain the purpose",
        r"brief overview of what the system does",
        r"enumerate system functions",
    )
    for pattern in banned_patterns:
        assert not re.search(pattern, full_text, flags=re.IGNORECASE), pattern


def test_content_states_the_implemented_arraylist_rules_without_unsupported_grading(artifact_path: Path) -> None:
    full_text = "\n".join(_paragraph_texts(_document(artifact_path)))
    required_facts = (
        "ArrayList<Student>",
        "Assessment 1",
        "Assessment 2",
        "Assessment 3",
        "0 through 100",
        "two decimal places",
        "tied",
        "in memory",
        "restarts",
    )
    for fact in required_facts:
        assert fact.lower() in full_text.lower(), fact

    unsupported_grading = (
        "passing grade",
        "failing grade",
        "weighted grade",
        "letter grade",
        "grade equivalent",
    )
    for claim in unsupported_grading:
        assert claim not in full_text.lower()


def test_system_features_section_has_exactly_eight_implemented_features(artifact_path: Path) -> None:
    doc = _document(artifact_path)
    paragraphs = doc.paragraphs
    start = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "2.3 System Features")
    feature_paragraphs = [p for p in paragraphs[start + 1 :] if p.text.strip()]

    assert len(feature_paragraphs) == 8
    assert all(p.style.name == "List Number" for p in feature_paragraphs)

    feature_text = "\n".join(p.text for p in feature_paragraphs).lower()
    for excluded in ("edit", "delete", "search", "sort", "rank", "export", "account", "database"):
        assert excluded not in feature_text


def test_letter_layout_margins_and_required_typography(artifact_path: Path) -> None:
    doc = _document(artifact_path)
    assert len(doc.sections) == 2
    for section in doc.sections:
        assert section.page_width.inches == pytest.approx(8.5, abs=0.01)
        assert section.page_height.inches == pytest.approx(11.0, abs=0.01)
        for margin in (section.top_margin, section.right_margin, section.bottom_margin, section.left_margin):
            assert margin.inches == pytest.approx(1.0, abs=0.01)

    normal = doc.styles["normal"]
    assert normal.font.name == "Arial"
    assert normal.font.size.pt == pytest.approx(11.0)

    heading_1 = doc.styles["Heading 1"]
    assert heading_1.font.name == "Arial"
    assert heading_1.font.size.pt == pytest.approx(17.0)
    assert heading_1.font.bold is True
    assert heading_1.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER

    heading_2 = doc.styles["Heading 2"]
    assert heading_2.font.name == "Arial"
    assert heading_2.font.size.pt == pytest.approx(13.0)
    assert heading_2.font.bold is True
    assert heading_2.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER

    body = _paragraph_after(doc, "1.1 Background of the Study")
    assert body.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert body.paragraph_format.line_spacing == pytest.approx(1.15)
    assert body.paragraph_format.first_line_indent.inches == pytest.approx(0.5, abs=0.01)


def test_cover_is_unnumbered_and_content_page_field_restarts_at_one(artifact_path: Path) -> None:
    doc = _document(artifact_path)
    parts = _xml_parts(artifact_path)

    assert len(doc.sections) == 2
    assert doc.sections[0].footer.is_linked_to_previous is False
    assert doc.sections[1].footer.is_linked_to_previous is False

    footer_parts = {
        name: content.decode("utf-8")
        for name, content in parts.items()
        if re.fullmatch(r"word/footer\d+\.xml", name)
    }
    assert len(footer_parts) == 2
    assert sum("PAGE" in xml for xml in footer_parts.values()) == 1
    numbered_footer = next(xml for xml in footer_parts.values() if "PAGE" in xml)
    assert 'w:jc w:val="center"' in numbered_footer

    document_xml = parts["word/document.xml"].decode("utf-8")
    section_properties = re.findall(r"<w:sectPr[\s\S]*?</w:sectPr>", document_xml)
    assert len(section_properties) == 2
    assert "w:pgNumType" not in section_properties[0]
    assert re.search(r'<w:pgNumType[^>]*w:fmt="decimal"[^>]*w:start="1"', section_properties[1])


def test_logo_and_flowchart_are_inline_images_with_meaningful_alt_text(artifact_path: Path) -> None:
    doc = _document(artifact_path)
    assert len(doc.inline_shapes) == 2

    parts = _xml_parts(artifact_path)
    document_xml = parts["word/document.xml"].decode("utf-8")
    descriptions = re.findall(r'<wp:docPr[^>]*descr="([^"]+)"', document_xml)
    assert descriptions == (
        ["Seal of St. John Paul II College of Davao.", FLOWCHART_ALT_TEXT]
        if descriptions
        else descriptions
    )
    assert len(descriptions) == 2
    assert FLOWCHART_ALT_TEXT in descriptions
    assert any("St. John Paul II College of Davao" in text for text in descriptions)
    assert "<wp:anchor" not in document_xml


def test_static_toc_and_caption_match_document_scope(artifact_path: Path) -> None:
    doc = _document(artifact_path)
    texts = _paragraph_texts(doc)
    toc_index = texts.index("TABLE OF CONTENTS")
    chapter_1_index = texts.index("CHAPTER I INTRODUCTION")
    toc_entries = texts[toc_index + 1 : chapter_1_index]

    assert toc_entries == [
        "Chapter I Introduction\t2",
        "1.1 Background of the Study\t2",
        "1.2 Objectives of the Study\t2",
        "1.3 Significance of the Study\t2",
        "Chapter II System Overview and Design\t3",
        "2.1 System Description\t3",
        "2.2 System Architecture and Flowchart\t3",
        "2.3 System Features\t5",
    ]
    assert texts.count("Figure 1 Student Score Management System Flowchart") == 1


def test_document_has_no_decorative_paragraph_rules(artifact_path: Path) -> None:
    parts = _xml_parts(artifact_path)
    document_xml = parts["word/document.xml"].decode("utf-8")
    assert "<w:pBdr" not in document_xml
